#!/usr/bin/env python3
"""Build the review sheet for GB 12268 names shared by several substances.

The regulatory import cannot pick a substance for a GB 12268 entry whose Chinese name matches more
than one substance (e.g. 甲醇 → 67-56-1 and a duplicate registry number); it leaves the entry
unlinked and opens a ``transport_unlinked`` review item. This sheet lists those entries with every
candidate substance and the assistant's draft suggestion (``link_suggestions.json``), for an EHS /
regulatory reviewer to decide. ``apply_link_sheet.py`` turns the filled sheet into
``app/resources/regulatory/gb12268_links.json``, which the import then follows.

用法:
    .venv/bin/python scripts/gb12268/build_link_sheet.py --out gb12268-关联核对表.xlsx
    .venv/bin/python scripts/gb12268/build_link_sheet.py --db-url sqlite:///./data/formunex.db --out x.xlsx
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from sqlalchemy import create_engine, select  # noqa: E402

from app.db.models_regulatory import RegulatoryReviewItem, SubstanceRow, SubstanceTransportRow  # noqa: E402

HERE = Path(__file__).resolve().parent
ALL = "全部候选"
NONE = "不关联"
SHEET = "核对表"
FONT = "Arial"
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [  # (header, width)
    ("序号", 6), ("UN 号", 8), ("中文名称", 16), ("英文名称", 30), ("类别", 7), ("包装类别", 9),
    ("候选物质（CAS｜英文名｜分子式｜危化目录序号｜IECSC）", 62), ("建议选定", 24), ("建议理由", 48),
    ("选定（填写）", 24), ("确认人（填写）", 12), ("确认日期（填写）", 14), ("备注（填写）", 24), ("状态", 9),
    ("候选 CAS（勿改）", 28),
]
COL = {h: i + 1 for i, (h, _) in enumerate(COLUMNS)}


def collect(db_url: str) -> list[dict]:
    engine = create_engine(db_url, future=True)
    out = []
    with engine.connect() as c:
        items = c.execute(select(RegulatoryReviewItem.name, RegulatoryReviewItem.detail).where(
            RegulatoryReviewItem.kind == "transport_unlinked", RegulatoryReviewItem.status == "open")).all()
        for name, detail in sorted(items, key=lambda r: (str((r.detail or {}).get("un_number")), r.name)):
            un = str((detail or {}).get("un_number") or "")
            rows = c.execute(select(SubstanceTransportRow.proper_shipping_name_en, SubstanceTransportRow.transport_class,
                                    SubstanceTransportRow.packing_group).where(
                SubstanceTransportRow.un_number == un, SubstanceTransportRow.proper_shipping_name_zh == name)).all()
            subs = c.execute(select(SubstanceRow.cas_no, SubstanceRow.name_en, SubstanceRow.mf,
                                    SubstanceRow.hazardous_catalog_no, SubstanceRow.iecsc_listed)
                             .where(SubstanceRow.name_zh == name).order_by(SubstanceRow.cas_no)).all()
            out.append({
                "un": un, "name_cn": name,
                "name_en": rows[0].proper_shipping_name_en if rows else "",
                "class": rows[0].transport_class if rows else "",
                "pg": "、".join(sorted({r.packing_group for r in rows if r.packing_group})) or "—",
                "candidates": [dict(cas=s.cas_no, name_en=s.name_en or "", mf=s.mf or "",
                                    catalog=s.hazardous_catalog_no, iecsc=s.iecsc_listed) for s in subs if s.cas_no],
            })
    engine.dispose()
    return out


def candidate_line(s: dict) -> str:
    catalog = f"目录 {s['catalog']}" if s["catalog"] is not None else "不在目录"
    iecsc = "IECSC 已列入" if s["iecsc"] else "IECSC 未查到"
    return f"{s['cas']}｜{s['name_en'] or '—'}｜{s['mf'] or '—'}｜{catalog}｜{iecsc}"


def build(entries: list[dict], suggestions: dict, out: Path) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = "说明"
    ws = wb.create_sheet(SHEET)

    # ------------------------------------------------------------------ 核对表
    for i, (head, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=i, value=head)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "C2"
    for n, e in enumerate(entries, 1):
        r = n + 1
        cas_list = [s["cas"] for s in e["candidates"]]
        sug = suggestions.get((e["un"], e["name_cn"]))
        sug_text = "—"
        if sug:
            sug_text = ALL if sorted(sug["cas"]) == sorted(cas_list) else "；".join(sug["cas"])
        values = {
            "序号": n, "UN 号": e["un"], "中文名称": e["name_cn"], "英文名称": e["name_en"], "类别": e["class"],
            "包装类别": e["pg"], "候选物质（CAS｜英文名｜分子式｜危化目录序号｜IECSC）": "\n".join(candidate_line(s) for s in e["candidates"]),
            "建议选定": sug_text, "建议理由": sug["reason"] if sug else "无建议，请人工判断",
            "状态": f'=IF({ws.cell(row=r, column=COL["选定（填写）"]).coordinate}="","待确认","已确认")',
            "候选 CAS（勿改）": "；".join(cas_list),
        }
        for head, v in values.items():
            ws.cell(row=r, column=COL[head], value=v)
        for head, _ in COLUMNS:
            cell = ws.cell(row=r, column=COL[head])
            cell.font = Font(name=FONT, color="808080" if head == "候选 CAS（勿改）" else None)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BOX
            if "（填写）" in head:
                cell.fill = INPUT_FILL
        ws.row_dimensions[r].height = max(30, 16 * len(e["candidates"]) + 6)
        options = [*cas_list, ALL, NONE]
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                            showErrorMessage=True, errorStyle="warning",
                            errorTitle="不在下拉选项中",
                            error="可直接填写多个 CAS（用“；”分隔）；请确认它们确实适用于该 UN 条目。")
        ws.add_data_validation(dv)
        dv.add(ws.cell(row=r, column=COL["选定（填写）"]))
    last = len(entries) + 1
    ws.cell(row=2, column=COL["选定（填写）"]).comment = Comment(
        "下拉选择一个 CAS、“全部候选”或“不关联”；也可填写多个 CAS，用“；”分隔。", "FormuNex")
    ws.auto_filter.ref = f"A1:{ws.cell(row=last, column=len(COLUMNS)).coordinate}"

    # ------------------------------------------------------------------ 说明
    status = ws.cell(row=2, column=COL["状态"]).column_letter
    lines = [
        ("GB 12268-2025 品名与物质关联核对表", True),
        (f"生成日期：{date.today().isoformat()}　　条目数：{len(entries)}", False),
        ("", False),
        ("用途", True),
        ("这些 GB 12268 条目的中文名称在物质库中对应多个物质（例如重复的 CAS 登记号、水合物、错挂的名称），"
         "系统无法自动判断该关联哪一个，暂未关联，运输分类无法直接引用。请逐条选定。", False),
        ("", False),
        ("填写方法", True),
        ("只填写黄色单元格：“选定”“确认人”“确认日期”“备注”。", False),
        ("“选定”下拉可选：某一个 CAS；“全部候选”（该 UN 条目同时适用于全部候选物质，例如不同水合物或复数条目）；"
         "“不关联”（候选都不适用）。也可直接填写多个 CAS，用“；”分隔。", False),
        ("“建议选定”“建议理由”由助手根据法规库数据整理，仅供参考，以 EHS / 法规专员的判断为准。", False),
        ("示例：UN 1230 甲醇 → 选定 67-56-1，确认人 张工，确认日期 2026-10-30。", False),
        ("", False),
        ("进度", True),
        ("已确认条目数", False),
        ("待确认条目数", False),
        ("", False),
        ("填完之后", True),
        ("把文件交回开发，运行 scripts/gb12268/apply_link_sheet.py 写入 gb12268_links.json 并提交；"
         "重新导入法规库后即按选定结果关联，相应复核项关闭。", False),
    ]
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 100
    for i, (text, bold) in enumerate(lines, 1):
        c = guide.cell(row=i, column=1 if bold or not text or i in (14, 15) else 2, value=text or None)
        c.font = Font(name=FONT, bold=bold, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    guide.cell(row=14, column=2, value=f'=COUNTIF({SHEET}!{status}2:{status}{last},"已确认")').font = Font(name=FONT)
    guide.cell(row=15, column=2, value=f'=COUNTIF({SHEET}!{status}2:{status}{last},"待确认")').font = Font(name=FONT)
    for r in (5, 9, 18):
        guide.row_dimensions[r].height = 34
    guide.cell(row=14, column=1).comment = Comment(f"统计“{SHEET}”中“选定”已填写的行。", "FormuNex")
    wb.active = 1
    wb.save(out)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--db-url", help="SQLAlchemy URL (default: FORMUNEX_DB_URL / settings.db_url)")
    parser.add_argument("--suggestions", type=Path, default=HERE / "link_suggestions.json")
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    if not args.db_url:
        from app.config import get_settings
        args.db_url = get_settings().db_url
    entries = collect(args.db_url)
    sug = {(s["un"], s["name_cn"]): s for s in json.loads(args.suggestions.read_text(encoding="utf-8"))["suggestions"]}
    build(entries, sug, args.out)
    print(json.dumps({"entries": len(entries), "with_suggestion": sum((e["un"], e["name_cn"]) in sug for e in entries),
                      "out": str(args.out)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
