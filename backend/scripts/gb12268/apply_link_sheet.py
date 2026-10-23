#!/usr/bin/env python3
"""Turn the filled GB 12268 link review sheet into ``gb12268_links.json``.

Reads the “核对表” sheet written by ``build_link_sheet.py``: each row whose “选定（填写）” cell is
filled becomes a person-decided link (a CAS, several CAS, “全部候选”, or “不关联”). Earlier decisions
in the JSON for other entries are kept; a row decided again replaces its old decision. Rows still
empty are skipped. Every CAS is checked (check digit); one outside the row's candidates is accepted
with a warning, because the reviewer may know a better registry number.

用法:
    .venv/bin/python scripts/gb12268/apply_link_sheet.py gb12268-关联核对表.xlsx
    .venv/bin/python scripts/seed_regulatory_db.py      # re-import: the links take effect
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from openpyxl import load_workbook  # noqa: E402

from app.services.plm.regulatory import cas as cas_mod  # noqa: E402

OUT = Path(__file__).resolve().parents[2] / "app" / "resources" / "regulatory" / "gb12268_links.json"
ALL, NONE, SHEET = "全部候选", "不关联", "核对表"


def _text(v) -> str:
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return "" if v is None else str(v).strip()


def read_sheet(path: Path) -> tuple[list[dict], list[str]]:
    ws = load_workbook(path, data_only=True)[SHEET]
    head = {_text(c.value): i for i, c in enumerate(ws[1])}
    need = ["UN 号", "中文名称", "选定（填写）", "确认人（填写）", "确认日期（填写）", "备注（填写）", "候选 CAS（勿改）"]
    missing = [h for h in need if h not in head]
    if missing:
        raise SystemExit(f"“{SHEET}”缺少列：{missing}（是否为 build_link_sheet.py 生成的表？）")
    links, problems = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        un, name, chosen = _text(row[head["UN 号"]]), _text(row[head["中文名称"]]), _text(row[head["选定（填写）"]])
        if not un or not chosen:
            continue
        candidates = [c for c in re.split(r"[；;,，\s]+", _text(row[head["候选 CAS（勿改）"]])) if c]
        if chosen == ALL:
            cas = candidates
        elif chosen == NONE:
            cas = []
        else:
            cas = [c for c in re.split(r"[；;,，\s]+", chosen) if c]
            bad = [c for c in cas if not cas_mod.is_valid(c)]
            if bad:
                problems.append(f"UN {un} {name}：CAS 校验不通过 {bad}，该行未写入")
                continue
            extra = [c for c in cas if c not in candidates]
            if extra:
                print(f"注意：UN {un} {name} 选定了候选之外的 CAS {extra}（物质库中没有时导入会提示）", file=sys.stderr)
        links.append({"un": un, "name_cn": name, "cas": cas,
                      "decided_by": _text(row[head["确认人（填写）"]]), "decided_at": _text(row[head["确认日期（填写）"]]),
                      "note": _text(row[head["备注（填写）"]])})
    return links, problems


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("sheet", type=Path)
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args(argv)
    links, problems = read_sheet(args.sheet)
    for p in problems:
        print(p, file=sys.stderr)
    existing = json.loads(args.out.read_text(encoding="utf-8")).get("links", []) if args.out.exists() else []
    merged = {(link["un"], link["name_cn"]): link for link in existing}
    for link in links:
        merged[(link["un"], link["name_cn"])] = link
    doc = {"meta": {"standard": "GB 12268-2025",
                    "note": "人工确认的品名—物质关联（名称对应多个物质时使用）；由 scripts/gb12268/apply_link_sheet.py 从核对表生成。"
                            "cas 为空表示确认不关联。"},
           "links": sorted(merged.values(), key=lambda link: (link["un"], link["name_cn"]))}
    args.out.write_text(json.dumps(doc, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(json.dumps({"written": len(links), "total": len(merged), "problems": len(problems), "out": str(args.out)},
                     ensure_ascii=False))
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
