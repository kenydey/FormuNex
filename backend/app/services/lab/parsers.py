"""Instrument export parsing (方案 §5.3.1, M7): configurable CSV / Excel templates → measurement rows.

A template says where the sample number is and which columns are which
metrics. Two layouts:

* ``wide`` — one row per reading, one column per metric (colorimeters,
  viscometers, gloss meters): ``columns = {"L*": {"metric": "color_l"}, …}``;
* ``long`` — one row per value: ``metric_column`` / ``value_column`` /
  ``unit_column``.

Several readings of one sample are aggregated (``mean`` / ``last`` / ``all``).
CSV is decoded as UTF-8 (with or without BOM) or GB18030, the delimiter
sniffed; Excel via openpyxl. Built-in templates cover common exports; others
are passed with the request or saved per instrument (``instruments.parser``).
"""
from __future__ import annotations

import csv
import io
import statistics
from dataclasses import dataclass
from typing import Any

from .methods import LabError

BUILTIN_TEMPLATES: dict[str, dict[str, Any]] = {
    "generic_long": {
        "label": "通用长表（样品号, 指标, 数值, 单位）", "layout": "long", "sample_column": "sample_no",
        "metric_column": "metric", "value_column": "value", "unit_column": "unit", "aggregate": "all",
    },
    "colorimeter_lab": {
        "label": "色差仪导出（Sample, L*, a*, b*, ΔE）", "layout": "wide", "sample_column": "Sample",
        "columns": {"L*": {"metric": "color_l"}, "a*": {"metric": "color_a"}, "b*": {"metric": "color_b"},
                    "ΔE": {"metric": "delta_e", "aliases": ["dE", "DE*ab", "ΔE*ab"]}},
        "aggregate": "mean",
    },
    "viscometer_rotational": {
        "label": "旋转黏度计导出（Sample ID, Viscosity (mPa·s), Temperature）", "layout": "wide",
        "sample_column": "Sample ID",
        "columns": {"Viscosity (mPa·s)": {"metric": "viscosity_mpas", "unit": "mPa·s",
                                          "aliases": ["Viscosity (cP)", "黏度(mPa·s)", "粘度(mPa·s)"]}},
        "aggregate": "last",
    },
    "gloss_meter": {
        "label": "光泽度仪导出（样品, 60°）", "layout": "wide", "sample_column": "样品",
        "columns": {"60°": {"metric": "gloss_60", "unit": "GU", "aliases": ["60", "Gloss 60°"]}},
        "aggregate": "mean",
    },
}


@dataclass
class Reading:
    sample_no: str
    metric: str
    value: float
    unit: str | None
    readings: int = 1


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise LabError("无法识别文件编码（支持 UTF-8 / GB18030）")


def _rows_csv(data: bytes, skip_rows: int) -> list[list[str]]:
    text = _decode(data)
    lines = text.splitlines()[skip_rows:]
    sample = "\n".join(lines[:20])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [r for r in csv.reader(lines, dialect) if any(c.strip() for c in r)]


def _rows_xlsx(data: bytes, skip_rows: int, sheet: str | None) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — any unreadable workbook is a user error
        raise LabError(f"无法读取 Excel 文件：{exc}") from None
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    return [r for r in rows[skip_rows:] if any(c.strip() for c in r)]


def _num(text: str, decimal_comma: bool = False) -> float | None:
    """Numbers as exported; "1,35" is 1.35 only when the template says the instrument writes decimal commas
    (otherwise "1,350" could be read either way)."""
    t = text.strip()
    t = t.replace(".", "").replace(",", ".") if decimal_comma else t.replace(",", "")
    try:
        return float(t)
    except ValueError:
        return None


def parse(data: bytes, template: dict[str, Any], *, filename: str = "") -> tuple[list[Reading], list[str]]:
    fmt = template.get("format") or ("xlsx" if filename.lower().endswith((".xlsx", ".xlsm")) else "csv")
    skip = int(template.get("skip_rows") or 0)
    rows = _rows_xlsx(data, skip, template.get("sheet")) if fmt == "xlsx" else _rows_csv(data, skip)
    if len(rows) < 2:
        raise LabError("文件中没有数据行")
    header = [h.strip() for h in rows[0]]
    idx = {h: i for i, h in enumerate(header)}

    def col(name: str, aliases: list[str] | None = None, required: bool = True) -> int | None:
        for n in [name, *(aliases or [])]:
            if n in idx:
                return idx[n]
        if required:
            raise LabError(f"文件缺少列“{name}”（表头：{', '.join(header)}）")
        return None

    sc = col(template.get("sample_column") or "sample_no")
    dc = bool(template.get("decimal_comma"))
    warnings: list[str] = []
    raw: list[tuple[str, str, float, str | None]] = []
    layout = template.get("layout") or "wide"
    if layout == "long":
        mc, vc = col(template.get("metric_column") or "metric"), col(template.get("value_column") or "value")
        uc = col(template.get("unit_column") or "unit", required=False)
        for n, r in enumerate(rows[1:], start=2):
            r = r + [""] * (len(header) - len(r))
            v = _num(r[vc], dc)
            if v is None:
                warnings.append(f"第 {n} 行数值无法识别：{r[vc]!r}")
                continue
            raw.append((r[sc].strip(), r[mc].strip(), v, r[uc].strip() if uc is not None else None))
    else:
        specs = template.get("columns") or {}
        if not specs:
            raise LabError("宽表模板需要 columns（列名 → 指标）")
        cols = []
        for name, spec in specs.items():
            ci = col(name, spec.get("aliases"), required=False)
            if ci is None:
                warnings.append(f"文件中没有列“{name}”，已跳过")
                continue
            cols.append((ci, spec["metric"], spec.get("unit")))
        if not cols:
            raise LabError("文件中没有模板定义的任何数据列")
        for n, r in enumerate(rows[1:], start=2):
            r = r + [""] * (len(header) - len(r))
            for ci, metric, unit in cols:
                if not r[ci].strip():
                    continue
                v = _num(r[ci], dc)
                if v is None:
                    warnings.append(f"第 {n} 行“{header[ci]}”数值无法识别：{r[ci]!r}")
                    continue
                raw.append((r[sc].strip(), metric, v, unit))
    agg = template.get("aggregate") or "all"
    out: list[Reading] = []
    if agg == "all":
        out = [Reading(s, m, v, u) for s, m, v, u in raw if s]
    else:
        grouped: dict[tuple[str, str], list[tuple[float, str | None]]] = {}
        for s, m, v, u in raw:
            if s:
                grouped.setdefault((s, m), []).append((v, u))
        for (s, m), vals in grouped.items():
            v = statistics.fmean(x for x, _u in vals) if agg == "mean" else vals[-1][0]
            out.append(Reading(s, m, v, vals[-1][1], len(vals)))
    if any(not s for s, *_ in raw):
        warnings.append("部分行缺少样品号，已跳过")
    return out, warnings
