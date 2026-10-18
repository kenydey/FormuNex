"""Supplier SDS import (补齐计划 1-5).

A supplier SDS file (PDF / DOCX / plain text) is parsed into a structured
record — product name, CAS, GHS classifications, H/P statements, supplier
info — and written into the substance database with ``source='supplier_sds'``
and confidence 0.95 (see ``source_priority.SOURCE_CONFIDENCE``).

Deliberate limits, all surfaced rather than guessed:

* **Single-substance SDS only.** A mixture SDS lists several CAS numbers in
  section 3; merging them into one substance row would be wrong, so the
  import refuses with an explicit error naming every CAS found.
* **Never guesses.** Whatever cannot be attributed to a known field lands in
  ``unparsed`` / ``review_reasons`` and is returned to the caller. A missing
  product name, CAS or classification is a parse error, not a silent skip.
* **Replace semantics per source.** A new SDS from a supplier supersedes the
  previous ``supplier_sds`` rows for the same substance (unlocked rows are
  replaced; locked rows — EHS-confirmed — are never touched, nor are rows
  from any other source).
* **No invented content.** Tests use synthetic SDS text written for the
  tests, never a real supplier's document.
"""
from __future__ import annotations

import hashlib
import re
import uuid
from dataclasses import dataclass, field

from ....db.models_regulatory import (
    RegulatoryDataRelease,
    SubstanceClassificationRow,
    SubstanceRow,
)
from . import cas as cas_mod
from .catalog_text_parser import parse_hazard_text as parse_catalog_text
from .hazard_classes import BY_CODE, english_class_to_code, h_code_for

SOURCE = "supplier_sds"
CONFIDENCE = 0.95

_CJK_RE = re.compile(r"[一-鿿]")

# GHS SDS section headers: "第2部分 危险性概述", "2. 危险性概述",
# "SECTION 2: Hazards identification", "2 Hazards identification".
# A line is a header only when its number AND its title keyword agree —
# a bare "3.5% 溶液" content line must never split a section.
_SECTION_HEAD_RE = re.compile(
    r"^\s*(?:第\s*(?P<num_zh>\d{1,2})\s*部分|SECTION\s+(?P<num_en>\d{1,2})\s*:?|(?P<num_bare>\d{1,2})[.\s])\s*"
    r"(?P<title>.{0,60})$",
    re.IGNORECASE | re.MULTILINE,
)
_SECTION_KEYWORDS: dict[int, tuple[str, ...]] = {
    1: ("化学品及企业标识", "企业标识", "identification"),
    2: ("危险性概述", "hazards identification", "hazard identification"),
    3: ("成分", "组成", "composition"),
    4: ("急救", "first aid"),
    5: ("消防", "firefighting"),
    6: ("泄漏", "accidental release"),
    7: ("操作处置", "handling and storage"),
    8: ("接触控制", "个体防护", "exposure controls"),
    9: ("理化特性", "physical and chemical properties"),
    10: ("稳定性", "反应性", "stability and reactivity"),
    11: ("毒理学", "toxicological"),
    12: ("生态学", "ecological"),
    13: ("废弃处置", "disposal"),
    14: ("运输", "transport"),
    15: ("法规", "regulatory"),
    16: ("其他信息", "other information"),
}

_H_CODE_RE = re.compile(r"\bH\d{3}(?:\s*\+\s*H\d{3})*\b")
_P_CODE_RE = re.compile(r"\bP\d{3}(?:\s*\+\s*P\d{3})*\b")
_CAS_TOKEN_RE = re.compile(r"\b\d{2,7}-\d{2}-\d\b")
_CAS_LABELED_RE = re.compile(
    r"CAS(?:\s*(?:号|No\.?|RN|登记号))?\s*[:：#]?\s*(\d{2,7}-\d{2}-\d)", re.IGNORECASE
)
_EN_CLASS_RE = re.compile(
    r"(?P<name>[A-Z][A-Za-z ,/()\-]{3,80}?)\s*\(?\s*"
    r"(?:Category\s+(?P<cat>\d[A-C]?)|Division\s+(?P<div>\d\.\d)|Type\s+(?P<type>[A-G])|(?P<unstable>Unstable explosive))"
    r"\)?",
    re.IGNORECASE,
)
_NOT_CLASSIFIED_RE = re.compile(
    r"(未分类|不属于危险品|不属于危险化学品|not classified|not a hazardous|no hazard class)", re.IGNORECASE
)
_PHONE_RE = re.compile(r"[\+\(]?\d[\d\-\(\) ]{6,}\d")


def _fold_fullwidth(text: str) -> str:
    return text.translate(str.maketrans("０１２３４５６７８９，：；（）［］＊－", "0123456789,:;()[]*-"))


@dataclass
class ParsedClassification:
    hazard_class: str
    category: str
    h_code: str | None = None
    qualifier: str = ""


@dataclass
class ParsedSds:
    product_name: str | None = None
    cas_no: str | None = None
    supplier_name: str | None = None
    supplier_address: str | None = None
    emergency_phone: str | None = None
    classifications: list[ParsedClassification] = field(default_factory=list)
    h_codes: list[str] = field(default_factory=list)
    p_codes: list[str] = field(default_factory=list)
    not_classified: bool = False
    unparsed: list[str] = field(default_factory=list)
    review_reasons: list[str] = field(default_factory=list)


class SdsParseError(ValueError):
    """The upload cannot be turned into a substance record. Message is user-facing."""


def extract_text(content: bytes, filename: str) -> str:
    """Extract plain text from an SDS upload (PDF / DOCX / text)."""
    name = filename.lower()
    if name.endswith(".pdf"):
        try:
            import fitz  # PyMuPDF
        except ImportError as exc:
            raise SdsParseError("PDF 解析需要 PyMuPDF，未安装") from exc
        try:
            doc = fitz.open(stream=content, filetype="pdf")
            try:
                return "\n".join(page.get_text() for page in doc)
            finally:
                doc.close()
        except SdsParseError:
            raise
        except Exception as exc:
            raise SdsParseError(f"PDF 解析失败：{exc}") from exc
    if name.endswith(".docx"):
        try:
            from io import BytesIO

            import docx
        except ImportError as exc:
            raise SdsParseError("DOCX 解析需要 python-docx，未安装") from exc
        try:
            document = docx.Document(BytesIO(content))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            return "\n".join(parts)
        except Exception as exc:
            raise SdsParseError(f"DOCX 解析失败：{exc}") from exc
    if name.endswith((".txt", ".md", ".text")):
        return content.decode("utf-8", errors="replace")
    raise SdsParseError(f"不支持的文件类型：{filename}（仅支持 .pdf / .docx / .txt）")


def _split_sections(text: str) -> dict[int, str]:
    """Split SDS text into its 16 GHS sections. Returns {section_no: body}."""
    matches = list(_SECTION_HEAD_RE.finditer(text))
    headers: list[tuple[int, int, int]] = []  # (section_no, header_end, match_start)
    for m in matches:
        num = m.group("num_zh") or m.group("num_en") or m.group("num_bare")
        try:
            no = int(num)
        except (TypeError, ValueError):
            continue
        if not 1 <= no <= 16:
            continue
        title = (m.group("title") or "").lower()
        if not any(kw.lower() in title for kw in _SECTION_KEYWORDS[no]):
            continue
        headers.append((no, m.end(), m.start()))
    sections: dict[int, str] = {}
    for i, (no, end, start) in enumerate(headers):
        stop = headers[i + 1][2] if i + 1 < len(headers) else len(text)
        sections[no] = text[end:stop]
    return sections


def _first_label_value(lines: list[str], labels: tuple[str, ...]) -> str | None:
    """Value after a `Label: value` line, or the next non-empty line when the
    label stands alone. Returns None when the label is absent."""
    for i, line in enumerate(lines):
        stripped = line.strip()
        for label in labels:
            if stripped.startswith(label):
                rest = stripped[len(label):].lstrip(":： ").strip()
                if rest:
                    return rest
                for nxt in lines[i + 1:]:
                    if nxt.strip():
                        return nxt.strip()
                return None
    return None


def _find_cas(sections: dict[int, str], full_text: str) -> tuple[str | None, list[str]]:
    """The single CAS for this SDS, or None. Returns (cas, all_valid_cas).

    Labelled ``CAS: x`` hits are authoritative. Otherwise scan section 3
    (composition) first, then the whole document; tokens preceded by ``EC``
    are EC numbers, not CAS, and are skipped. Every candidate must pass the
    CAS checksum.
    """
    labeled = [c for c in _CAS_LABELED_RE.findall(full_text) if cas_mod.normalize(c)]
    if labeled:
        distinct = sorted({cas_mod.normalize(c) for c in labeled if cas_mod.normalize(c)})
        return distinct[0] if len(distinct) == 1 else None, distinct

    def scan(blob: str) -> list[str]:
        found: list[str] = []
        for m in _CAS_TOKEN_RE.finditer(blob):
            if re.search(r"\bEC\b", blob[max(0, m.start() - 6):m.start()], re.IGNORECASE):
                continue
            norm = cas_mod.normalize(m.group(0))
            if norm and norm not in found:
                found.append(norm)
        return found

    candidates = scan(sections.get(3, ""))
    if not candidates:
        candidates = scan(full_text)
    if len(candidates) == 1:
        return candidates[0], candidates
    return None, candidates


def _parse_section2_classifications(text: str) -> tuple[list[ParsedClassification], list[str]]:
    """Classifications from SDS section 2 (hazard identification).

    Chinese text goes through the catalogue parser (same
    ``急性毒性-经口,类别 3`` shape); English ``Acute toxicity, oral,
    Category 3`` lines are mapped via :func:`english_class_to_code`.
    """
    out: list[ParsedClassification] = []
    seen: set[tuple[str, str]] = set()

    def add(hc: str, cat: str, h: str | None, qualifier: str = "") -> None:
        key = (hc, cat)
        if key in seen or hc not in BY_CODE:
            return
        seen.add(key)
        out.append(ParsedClassification(
            hazard_class=hc, category=cat,
            h_code=h or h_code_for(hc, cat), qualifier=qualifier,
        ))

    result = parse_catalog_text(text)
    for item in result.items:
        add(item.hazard_class, item.category, item.h_code, item.qualifier)
    unparsed = list(result.unparsed)

    for m in _EN_CLASS_RE.finditer(text):
        code = english_class_to_code(m.group("name") or "")
        if not code:
            continue
        if m.group("unstable"):
            add(code, "unstable", None)
        elif m.group("div"):
            add(code, m.group("div"), None)
        elif m.group("type"):
            add(code, m.group("type"), None)
        elif m.group("cat"):
            add(code, m.group("cat"), None)
    return out, unparsed


def parse_sds_text(text: str) -> ParsedSds:
    """Parse extracted SDS text. Raises :class:`SdsParseError` when a required
    field (product name, CAS, classification) is missing or ambiguous."""
    if not text or not text.strip():
        raise SdsParseError("文件内容为空，无法解析")
    folded = _fold_fullwidth(text)
    sections = _split_sections(folded)
    sec1 = sections.get(1, folded)
    sec2 = sections.get(2, "")
    lines = [ln for ln in sec1.splitlines()]

    parsed = ParsedSds()
    parsed.product_name = _first_label_value(
        lines, ("产品名称", "化学品名称", "商品名", "Product name", "Product identifier", "Trade name")
    )
    parsed.supplier_name = _first_label_value(
        lines, ("供应商", "供货商", "制造商", "生产商", "公司", "Supplier", "Manufacturer", "Company")
    )
    parsed.supplier_address = _first_label_value(lines, ("地址", "Address",))
    phone = _first_label_value(
        lines, ("应急电话", "应急咨询电话", "Emergency telephone", "Emergency phone", "应急电话号码")
    )
    if phone:
        m = _PHONE_RE.search(phone)
        parsed.emergency_phone = m.group(0).strip() if m else phone

    cas, all_cas = _find_cas(sections, folded)
    if all_cas and len(all_cas) > 1:
        raise SdsParseError(
            "该 SDS 含多个 CAS 号（混合物），本导入仅支持单组分物质 SDS："
            + "、".join(all_cas)
        )
    parsed.cas_no = cas

    if sec2:
        classifications, unparsed = _parse_section2_classifications(sec2)
        parsed.classifications = classifications
        parsed.unparsed = unparsed
        codes: list[str] = []
        for chunk in _H_CODE_RE.findall(sec2):
            for code in re.split(r"\s*\+\s*", chunk):
                if code and code not in codes:
                    codes.append(code)
        parsed.h_codes = codes
        pcodes: list[str] = []
        for chunk in _P_CODE_RE.findall(sec2):
            for code in re.split(r"\s*\+\s*", chunk):
                if code and code not in pcodes:
                    pcodes.append(code)
        parsed.p_codes = pcodes
        if not classifications and _NOT_CLASSIFIED_RE.search(sec2):
            parsed.not_classified = True
            parsed.review_reasons.append("SDS 第 2 部分声明未分类，未写入分类")
    else:
        parsed.review_reasons.append("未找到第 2 部分（危险性概述），未解析分类")

    errors: list[str] = []
    if not parsed.product_name:
        errors.append("未解析到产品名称（第 1 部分需含“产品名称”）")
    if not parsed.cas_no:
        errors.append("未解析到 CAS 号（第 3 部分或 CAS 标注行需含有效 CAS 号）")
    if not parsed.classifications and not parsed.not_classified:
        errors.append("未解析到 GHS 分类（第 2 部分需含“危害类别,类别 N”或英文 Category 表述）")
    if errors:
        raise SdsParseError("；".join(errors))
    if not parsed.supplier_name:
        parsed.review_reasons.append("未解析到供应商名称")
    return parsed


def file_digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def import_sds(session, parsed: ParsedSds, *, filename: str, digest: str) -> dict:
    """Write a parsed SDS into the substance database.

    Flushes only — the caller commits. Replaces this substance's unlocked
    ``supplier_sds`` classifications; locked rows and other sources are kept.
    """
    cas = parsed.cas_no or ""
    substance = (
        session.query(SubstanceRow).filter(SubstanceRow.cas_no == cas).first()
        if cas
        else None
    )
    if substance is None:
        name = parsed.product_name or ""
        is_zh = bool(_CJK_RE.search(name))
        substance = SubstanceRow(
            id=str(uuid.uuid4()),
            cas_no=cas or None,
            name_zh=name if is_zh else "",
            name_en="" if is_zh else name,
            aliases=[],
        )
        session.add(substance)
        session.flush()
    elif parsed.product_name:
        # Fill in a missing display name, never overwrite an existing one.
        if _CJK_RE.search(parsed.product_name) and not substance.name_zh:
            substance.name_zh = parsed.product_name
        elif not _CJK_RE.search(parsed.product_name) and not substance.name_en:
            substance.name_en = parsed.product_name

    release = RegulatoryDataRelease(
        id=str(uuid.uuid4()),
        label=f"sds-import:{filename}",
        source="supplier_sds_import",
        stats={
            "filename": filename,
            "cas_no": cas,
            "classifications": len(parsed.classifications),
            "h_codes": parsed.h_codes,
            "p_codes": parsed.p_codes,
            "supplier": parsed.supplier_name or "",
        },
        input_digest=digest,
    )
    session.add(release)

    replaced = (
        session.query(SubstanceClassificationRow)
        .filter(
            SubstanceClassificationRow.substance_id == substance.id,
            SubstanceClassificationRow.source == SOURCE,
            SubstanceClassificationRow.locked_at.is_(None),
        )
        .delete(synchronize_session=False)
    )
    for item in parsed.classifications:
        session.add(SubstanceClassificationRow(
            id=str(uuid.uuid4()),
            substance_id=substance.id,
            hazard_class=item.hazard_class,
            category=item.category,
            h_code=item.h_code,
            qualifier=item.qualifier,
            source=SOURCE,
            source_ref=f"sds-import:{filename}",
            confidence=CONFIDENCE,
            release_id=release.id,
        ))
    session.flush()
    return {
        "substance_id": substance.id,
        "cas_no": cas,
        "product_name": parsed.product_name,
        "supplier_name": parsed.supplier_name,
        "classifications_written": len(parsed.classifications),
        "classifications_replaced": replaced,
        "h_codes": parsed.h_codes,
        "p_codes": parsed.p_codes,
        "not_classified": parsed.not_classified,
        "unparsed": parsed.unparsed,
        "review_reasons": parsed.review_reasons,
        "release_id": release.id,
    }
