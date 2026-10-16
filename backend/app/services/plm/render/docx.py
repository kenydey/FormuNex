"""Render structured document content to Word.

With a company template (``.docx`` with docxtpl / Jinja tags) the content is
passed as context — ``{{ header.product_name }}``, ``{% for s in sections %}``
— so layout is owned by whoever edits the template in Word. Without one, a
plain default layout is generated with python-docx.
"""
from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def _text(value: Any) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, list):
        return "\n".join(_text(v) for v in value)
    if isinstance(value, dict):
        if "code" in value and "text" in value:
            return f"{value['code']} {value['text']}"
        return "；".join(f"{k}: {v}" for k, v in value.items())
    return str(value)


def _set_cjk_font(doc, name: str = "SimSun") -> None:
    style = doc.styles["Normal"]
    style.font.size = Pt(10)
    style.font.name = name  # creates w:rFonts
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def render_docx(content: dict, template: bytes | None = None) -> bytes:
    if template is not None:
        from docxtpl import DocxTemplate

        tpl = DocxTemplate(io.BytesIO(template))
        tpl.render({**content, "text": _text})
        out = io.BytesIO()
        tpl.save(out)
        return out.getvalue()

    doc = Document()
    _set_cjk_font(doc)
    header = content.get("header") or {}
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = " · ".join(
        str(b) for b in (header.get("product_name"), content.get("title"), header.get("version")) if b)
    h = doc.add_heading(content.get("title", ""), level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = " · ".join(str(b) for b in (header.get("product_name"), header.get("product_code"),
                                        header.get("revision_date"), header.get("standard")) if b)
    if meta:
        p = doc.add_paragraph(meta)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for s in content.get("sections", []):
        doc.add_heading(f"{s['no']}  {s['title']}", level=1)
        if s.get("fields"):
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for f in s["fields"]:
                cells = table.add_row().cells
                cells[0].text = f.get("label") or ""
                cells[1].text = _text(f.get("value"))
        for tb in s.get("tables", []):
            cols = tb.get("columns") or []
            table = doc.add_table(rows=1, cols=max(1, len(cols)))
            table.style = "Table Grid"
            for i, c in enumerate(cols):
                table.rows[0].cells[i].text = str(c)
            for row in tb.get("rows", []):
                cells = table.add_row().cells
                for i, v in enumerate(row[: len(cols)]):
                    cells[i].text = _text(v)
        for para in s.get("paragraphs", []):
            doc.add_paragraph(para)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
