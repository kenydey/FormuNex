"""GB 12268 link review sheet: build → fill → apply round trip."""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parents[1] / "scripts" / "gb12268"


def _load(name: str):
    spec = importlib.util.spec_from_file_location(name, HERE / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


build_mod = _load("build_link_sheet")
apply_mod = _load("apply_link_sheet")

ENTRIES = [
    {"un": "1230", "name_cn": "甲醇", "name_en": "METHANOL", "class": "3", "pg": "II", "candidates": [
        {"cas": "170082-17-4", "name_en": "METHANOL", "mf": "CH4O", "catalog": None, "iecsc": None},
        {"cas": "67-56-1", "name_en": "Methanol", "mf": "CH4O", "catalog": 1022, "iecsc": True}]},
    {"un": "1405", "name_cn": "硅化钙", "name_en": "CALCIUM SILICIDE", "class": "4.3", "pg": "II、III", "candidates": [
        {"cas": "12013-55-7", "name_en": "", "mf": "CaSi", "catalog": 839, "iecsc": True},
        {"cas": "12013-56-8", "name_en": "", "mf": "CaSi2", "catalog": 838, "iecsc": True}]},
    {"un": "3495", "name_cn": "碘", "name_en": "IODINE", "class": "8", "pg": "III", "candidates": [
        {"cas": "7553-56-2", "name_en": "Iodine", "mf": "I2", "catalog": None, "iecsc": True},
        {"cas": "12190-71-5", "name_en": "iodine", "mf": "I2", "catalog": None, "iecsc": True}]},
    {"un": "2647", "name_cn": "丙二腈", "name_en": "MALONONITRILE", "class": "6.1", "pg": "II", "candidates": [
        {"cas": "109-77-3", "name_en": "Malononitrile", "mf": "C3H2N2", "catalog": 115, "iecsc": True},
        {"cas": "69567-10-8", "name_en": "", "mf": "", "catalog": None, "iecsc": None}]},
]
SUGGESTIONS = {("1230", "甲醇"): {"cas": ["67-56-1"], "reason": "目录 1022"},
               ("1405", "硅化钙"): {"cas": ["12013-55-7", "12013-56-8"], "reason": "两者均适用"}}


def _fill(path: Path, values: dict[str, tuple[str, str]]) -> None:
    wb = load_workbook(path)
    ws = wb["核对表"]
    head = {c.value: c.column for c in ws[1]}
    for row in ws.iter_rows(min_row=2):
        un = row[head["UN 号"] - 1].value
        if un in values:
            chosen, who = values[un]
            ws.cell(row=row[0].row, column=head["选定（填写）"], value=chosen)
            ws.cell(row=row[0].row, column=head["确认人（填写）"], value=who)
    wb.save(path)


def test_sheet_lists_candidates_suggestions_and_input_cells(tmp_path):
    out = tmp_path / "s.xlsx"
    build_mod.build(ENTRIES, SUGGESTIONS, out)
    wb = load_workbook(out)
    ws = wb["核对表"]
    head = {c.value: c.column for c in ws[1]}
    rows = {r[head["UN 号"] - 1].value: r for r in ws.iter_rows(min_row=2)}
    assert rows["1230"][head["建议选定"] - 1].value == "67-56-1"
    assert rows["1405"][head["建议选定"] - 1].value == "全部候选"      # suggestion = every candidate
    assert rows["3495"][head["建议选定"] - 1].value == "—"             # no suggestion: never invented
    assert "目录 1022" in rows["1230"][head["候选物质（CAS｜英文名｜分子式｜危化目录序号｜IECSC）"] - 1].value
    assert rows["1230"][head["选定（填写）"] - 1].value is None          # left for the reviewer
    assert rows["1230"][head["选定（填写）"] - 1].fill.fgColor.rgb[-6:] == "FFFF00"  # yellow input cell
    assert len(ws.data_validations.dataValidation) == len(ENTRIES)
    assert wb.sheetnames == ["说明", "核对表"]


def test_filled_sheet_becomes_links_and_merges_with_earlier_decisions(tmp_path):
    sheet, out = tmp_path / "s.xlsx", tmp_path / "links.json"
    build_mod.build(ENTRIES, SUGGESTIONS, sheet)
    out.write_text(json.dumps({"links": [{"un": "2013", "name_cn": "磷化锶", "cas": ["12504-13-1"]},
                                         {"un": "1230", "name_cn": "甲醇", "cas": ["170082-17-4"]}]}), encoding="utf-8")
    _fill(sheet, {"1230": ("67-56-1", "张工"), "1405": ("全部候选", "张工"), "3495": ("7553-56-2；12190-71-5", "李工")})
    assert apply_mod.main([str(sheet), "--out", str(out)]) == 0
    links = {(x["un"], x["name_cn"]): x for x in json.loads(out.read_text(encoding="utf-8"))["links"]}
    assert links[("1230", "甲醇")]["cas"] == ["67-56-1"] and links[("1230", "甲醇")]["decided_by"] == "张工"
    assert links[("1405", "硅化钙")]["cas"] == ["12013-55-7", "12013-56-8"]
    assert links[("3495", "碘")]["cas"] == ["7553-56-2", "12190-71-5"]
    assert links[("2013", "磷化锶")]["cas"] == ["12504-13-1"]          # earlier decision kept
    assert ("2647", "丙二腈") not in links                             # still empty in the sheet: skipped


def test_not_linked_and_invalid_cas(tmp_path):
    sheet, out = tmp_path / "s.xlsx", tmp_path / "links.json"
    build_mod.build(ENTRIES, SUGGESTIONS, sheet)
    _fill(sheet, {"1230": ("不关联", "张工"), "2647": ("109-77-4", "张工")})  # wrong check digit
    assert apply_mod.main([str(sheet), "--out", str(out)]) == 1
    links = {(x["un"], x["name_cn"]): x for x in json.loads(out.read_text(encoding="utf-8"))["links"]}
    assert links[("1230", "甲醇")]["cas"] == []
    assert ("2647", "丙二腈") not in links
