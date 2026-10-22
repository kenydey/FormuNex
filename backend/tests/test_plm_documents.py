"""PLM documents: SDS / TDS / COA / label content, rendering, versioning, diff, API."""
from __future__ import annotations

import io
import uuid
from datetime import datetime
from decimal import Decimal

import pytest

pytest.importorskip("reportlab")
pytest.importorskip("docx")
pytest.importorskip("svglib")

from fastapi.testclient import TestClient  # noqa: E402

from app.config import get_settings  # noqa: E402
from app.services.plm.diff import diff_documents  # noqa: E402
from app.services.plm.documents.coa_builder import (  # noqa: E402
    CoaContext,
    ResultItem,
    SpecItem,
    build_coa,
    judge,
    round_gbt8170,
)
from app.services.plm.documents.confidential_ranges import band, disclose  # noqa: E402
from app.services.plm.documents.label_builder import LabelContext, build_label, check_label_size  # noqa: E402
from app.services.plm.documents.sds_builder import SECTION_TITLES, SdsContext, build_sds  # noqa: E402
from app.services.plm.documents.tds_builder import TdsContext, build_tds, split_list_field  # noqa: E402
from app.services.plm.ghs.engine import classify  # noqa: E402
from app.services.plm.ghs.model import Component, ComponentHazard, MixtureProperties  # noqa: E402

COMPANY = {"name": "方枢示例化工有限公司", "address": "上海市示例路 1 号", "phone": "021-0000", "emergency_phone": "021-0001"}
PRODUCT = {"name_zh": "FX-100 环氧底漆", "name_en": "FX-100 Epoxy Primer", "code": "FX-100"}


def _classification(corrosive=False):
    comps = [Component("二甲苯", 25, "1330-20-7", hazards=[ComponentHazard("flammable_liquid", "3", "H226"),
                                                         ComponentHazard("skin_corrosion", "2", "H315")]),
             Component("环氧树脂", 60, "25068-38-6", hazards=[ComponentHazard("skin_sensitizer", "1", "H317")]),
             Component("水", 15, "7732-18-5")]
    if corrosive:
        comps[2] = Component("氢氧化钠", 15, "1310-73-2", hazards=[ComponentHazard("skin_corrosion", "1A", "H314")])
    r = classify(comps, MixtureProperties(flash_point_c=27, initial_boiling_point_c=138)).to_dict()
    r["statements"] = {c: {"zh": f"{c}中文", "en": ""} for c in r["h_codes"] + r["p_codes"]}
    r["statements"]["P501"] = {"zh": "处置内装物/容器……", "en": ""}
    return r


def _sds(**kw):
    cls = kw.pop("cls", None) or _classification()
    ctx = SdsContext(product=PRODUCT, company=kw.pop("company", COMPANY), classification=cls,
                     statements=cls["statements"], components=kw.pop("components", cls["disclosure"]),
                     revision={"version": "A.1", "date": "2026-10-16"}, **kw)
    return build_sds(ctx)


# --------------------------------------------------------------------------- content builders


@pytest.mark.parametrize(("pct", "expected"), [(14.5, "10%～20%"), (0.05, "<0.1%"), (1, "1%～3%"), (100, "90%～100%")])
def test_trade_secret_bands(pct, expected):
    assert band(pct) == expected
    assert disclose(14.5, trade_secret=False) == "14.5%"


def test_sds_has_sixteen_gb_sections_in_order():
    sds = _sds()
    assert [s["no"] for s in sds["sections"]] == list(range(1, 17))
    assert [s["key"] for s in sds["sections"]] == [k for k, _, _ in SECTION_TITLES]
    assert sds["sections"][1]["title"] == "危险性概述"
    assert sds["validation"]["missing"] == []


def test_sds_hazard_section_comes_from_the_engine():
    sds = _sds()
    hz = {f["key"]: f for f in sds["sections"][1]["fields"]}
    assert hz["signal_word"]["value"] == "警告" and hz["signal_word"]["source"] == "computed"
    assert hz["hazard_statements"]["value"][0].startswith("H226")
    assert "P501: 防范说明含待补充内容" in "".join(sds["validation"]["warnings"])


def test_sds_trade_secret_composition_is_banded_with_generic_name():
    cls = _classification()
    comps = [dict(c, trade_secret=c["name"] == "环氧树脂", generic_name="改性环氧树脂" if c["name"] == "环氧树脂" else None)
             for c in cls["disclosure"]]
    sds = _sds(cls=cls, components=comps)
    rows = sds["sections"][2]["tables"][0]["rows"]
    resin = next(r for r in rows if r[0] == "改性环氧树脂")
    assert resin[2] == "60%～70%"
    xylene = next(r for r in rows if r[0] == "二甲苯")
    assert xylene[2] == "25%"


def test_sds_reports_missing_mandatory_fields():
    sds = _sds(company={"name": "x", "address": "y"})
    assert "1.24小时应急电话" in sds["validation"]["missing"]


def test_sds_first_aid_follows_the_hazard():
    mild = {f["key"]: f["value"] for f in _sds()["sections"][3]["fields"]}
    severe = {f["key"]: f["value"] for f in _sds(cls=_classification(corrosive=True))["sections"][3]["fields"]}
    assert "15 分钟" not in mild["skin"] and "15 分钟" in severe["skin"]
    assert "禁止催吐" in severe["ingestion"]


def test_sds_english_titles_and_missing_translation_warning():
    sds = _sds(lang="en-US")
    assert sds["title"] == "Safety Data Sheet" and sds["sections"][1]["title"] == "Hazards identification"
    assert any("English statement text missing" in w for w in sds["validation"]["warnings"])


def test_sds_section9_marks_unknown_properties():
    sds = _sds(properties={"flash_point_c": {"value": 27, "source": "measured", "method": "GB/T 261"}})
    rows = {r[0]: r for r in sds["sections"][8]["tables"][0]["rows"]}
    assert rows["闪点（闭杯）"][1:] == ["27 ℃", "GB/T 261", "measured"]
    assert rows["蒸气压"][1] == "无资料"


def test_tds_keeps_measurement_ids_for_cross_document_checks():
    tds = build_tds(TdsContext(product={**PRODUCT, "description": "双组分环氧底漆"}, company=COMPANY,
                               properties=[{"key": "density_g_cm3", "label": "密度", "value": 1.35, "unit": "g/cm³",
                                            "method": "GB/T 6750", "measurement_id": "m-1"}],
                               application={"method": "喷涂"}))
    assert tds["meta"]["measurement_ids"] == {"density_g_cm3": "m-1"}
    assert tds["validation"]["missing"] == []
    assert "产品说明" in build_tds(TdsContext(product=PRODUCT, company=COMPANY))["validation"]["missing"]


def test_gbt8170_rounding_is_half_to_even():
    assert round_gbt8170(1.05, 1) == Decimal("1.0")
    assert round_gbt8170(1.15, 1) == Decimal("1.2")
    assert round_gbt8170(1.051, 1) == Decimal("1.1")


@pytest.mark.parametrize(("spec", "result", "verdict"), [
    (SpecItem("visc", max=1.0), ResultItem("visc", 1.05), "FAIL"),                      # full-value comparison (default)
    (SpecItem("visc", max=1.0, comparison="rounded"), ResultItem("visc", 1.05), "PASS"),  # rounds to 1.0
    (SpecItem("visc", max=1.0, comparison="rounded"), ResultItem("visc", 1.15), "FAIL"),
    (SpecItem("solids", min=55, max=60), ResultItem("solids", 55), "PASS"),            # limits inclusive
    (SpecItem("look", text="均匀无结块"), ResultItem("look", value_text="均匀", conforms=True), "PASS"),
    (SpecItem("look", text="均匀无结块"), ResultItem("look", value_text="有结块", conforms=False), "FAIL"),
    (SpecItem("solids", min=55), None, "NO_DATA"),
])
def test_coa_judgement(spec, result, verdict):
    assert judge(spec, result)[0] == verdict


def test_coa_overall_conclusion():
    specs = [SpecItem("solids", "固含量", "%", "GB/T 1725", 55, 60), SpecItem("visc", "黏度", "s", "GB/T 1723", max=90)]
    ok = build_coa(CoaContext(PRODUCT, COMPANY, {"batch_no": "B1"}, specs,
                              [ResultItem("solids", 57.2), ResultItem("visc", 80)]))
    assert ok["meta"]["overall"] == "PASS" and ok["meta"]["issuable"]
    bad = build_coa(CoaContext(PRODUCT, COMPANY, {"batch_no": "B1"}, specs, [ResultItem("solids", 54.9)]))
    assert bad["meta"]["overall"] == "FAIL" and not bad["meta"]["issuable"]
    partial = build_coa(CoaContext(PRODUCT, COMPANY, {"batch_no": "B1"}, specs, [ResultItem("solids", 57)]))
    assert partial["meta"]["overall"] == "INCOMPLETE"


@pytest.mark.parametrize(("volume", "size", "ok"), [(0.05, (20, 20), True), (2, (50, 75), True), (2, (40, 75), False),
                                                     (20, (75, 100), True), (200, (75, 100), False), (2000, (200, 300), True)])
def test_gb15258_label_size(volume, size, ok):
    assert check_label_size(volume, *size)["ok"] is ok


def test_label_content():
    cls = _classification()
    lab = build_label(LabelContext(PRODUCT, COMPANY, cls, cls["statements"], ["二甲苯"], container_volume_l=200,
                                   size_mm=(75, 100)))
    f = {x["key"]: x["value"] for x in lab["sections"][0]["fields"]}
    assert f["pictograms"] == [p["code"] for p in cls["pictograms"]]
    assert [s["code"] for s in f["precautionary_statements"]] == cls["label_p_codes"]
    assert any("100×150" in w for w in lab["validation"]["warnings"])


def test_diff_flags_major_changes():
    a, b = _sds(), _sds(cls=_classification(corrosive=True))
    d = diff_documents(a, b)
    assert d["major"] and any(c["path"].startswith("hazards.") for c in d["changes"])
    minor = _sds()
    minor["sections"][0]["fields"][2]["value"] = "FX-100B"
    d2 = diff_documents(a, minor)
    assert d2["count"] == 1 and not d2["major"]
    assert diff_documents(a, a)["count"] == 0


# --------------------------------------------------------------------------- rendering


def test_pdf_render_contains_the_chinese_text():
    pypdf = pytest.importorskip("pypdf")
    from app.services.plm.render.pdf import render_pdf

    pdf, info = render_pdf(_sds())
    assert pdf.startswith(b"%PDF")
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) >= 3
    if info["font_embedded"]:  # the CID fallback has no ToUnicode text to extract
        text = "".join(p.extract_text() for p in reader.pages)
        assert "危险性概述" in text and "第 1 页" in text


def test_docx_default_and_template_rendering():
    import docx
    from docx import Document

    from app.services.plm.render.docx import render_docx

    sds = _sds()
    out = docx.Document(io.BytesIO(render_docx(sds)))
    assert any(p.text.startswith("2  危险性概述") for p in out.paragraphs)
    tpl = Document()
    tpl.add_paragraph("产品：{{ header.product_name }}")
    tpl.add_paragraph("{% for s in sections %}{{ s.no }}. {{ s.title }}；{% endfor %}")
    buf = io.BytesIO()
    tpl.save(buf)
    rendered = docx.Document(io.BytesIO(render_docx(sds, buf.getvalue())))
    texts = [p.text for p in rendered.paragraphs]
    assert texts[0] == "产品：FX-100 环氧底漆" and "16. 其他信息" in texts[1]


def test_label_pdf_has_the_physical_size_and_detects_overflow():
    pypdf = pytest.importorskip("pypdf")
    from reportlab.lib.units import mm

    from app.services.plm.render.label_pdf import render_label_pdf

    cls = _classification(corrosive=True)
    lab = build_label(LabelContext(PRODUCT, COMPANY, cls, cls["statements"], size_mm=(100, 150), qr_payload="x"))
    pdf, info = render_label_pdf(lab)
    box = pypdf.PdfReader(io.BytesIO(pdf)).pages[0].mediabox
    assert (round(float(box.width) / mm), round(float(box.height) / mm)) == (100, 150)
    assert info["overflow"] is False
    tiny = build_label(LabelContext(PRODUCT, COMPANY, cls, cls["statements"], size_mm=(30, 30)))
    assert render_label_pdf(tiny)[1]["overflow"] is True


# --------------------------------------------------------------------------- API end to end


@pytest.fixture()
def api(tmp_path, monkeypatch):
    monkeypatch.setenv("FORMUNEX_DB_URL", f"sqlite:///{tmp_path / 'docs.db'}")
    monkeypatch.setenv("FORMUNEX_API_AUTH_ENABLED", "false")
    monkeypatch.setenv("FORMUNEX_PLM_STORAGE_DIR", str(tmp_path / "plm"))
    monkeypatch.setenv("FORMUNEX_PLM_COMPANY_NAME", "方枢示例化工有限公司")
    monkeypatch.setenv("FORMUNEX_PLM_COMPANY_ADDRESS", "上海市示例路 1 号")
    monkeypatch.setenv("FORMUNEX_PLM_EMERGENCY_PHONE", "021-0001")
    get_settings.cache_clear()
    from app.db.database import default_session_factory
    from app.db.models import FormulationVersion
    from app.db.models_regulatory import SubstanceClassificationRow, SubstanceOELRow, SubstanceRow

    with default_session_factory()() as s:
        xyl = SubstanceRow(id=str(uuid.uuid4()), cas_no="1330-20-7", name_zh="二甲苯", name_en="xylene", aliases=[],
                           hazardous_catalog_no=2289)
        s.add(xyl)
        s.add(SubstanceRow(id=str(uuid.uuid4()), cas_no="7732-18-5", name_zh="水", name_en="water", aliases=[]))
        s.flush()
        for hc, cat, h in [("flammable_liquid", "3", "H226"), ("skin_corrosion", "2", "H315")]:
            s.add(SubstanceClassificationRow(id=str(uuid.uuid4()), substance_id=xyl.id, hazard_class=hc, category=cat,
                                             h_code=h, source="cn_catalog"))
        s.add(SubstanceOELRow(id=str(uuid.uuid4()), substance_id=xyl.id, pc_twa=50, pc_stel=100))
        s.add(FormulationVersion(id="fv-1", lineage_id="l", version=1, domain="coating", name="v1", created_at=datetime(2026, 10, 1),
                                 snapshot={"ingredients": [{"name": "二甲苯", "cas_no": "1330-20-7", "weight_pct": 30},
                                                           {"name": "水", "cas_no": "7732-18-5", "weight_pct": 70}]}))
        s.commit()
    from app.main import app

    client = TestClient(app)
    pid = client.post("/api/plm/products", json={"code": "FX-100", "name_zh": "FX-100 环氧底漆"}).json()["id"]
    rid = client.post(f"/api/plm/products/{pid}/revisions", json={
        "formulation_version_id": "fv-1",
        "spec_snapshot": {"release_specs": [{"metric": "solids", "label": "固含量", "unit": "%", "min": 55, "max": 60},
                                            {"metric": "visc", "label": "黏度", "unit": "s", "max": 90}]}}).json()["id"]
    yield client, pid, rid
    get_settings.cache_clear()


def test_sds_generation_versioning_and_files(api):
    client, pid, rid = api
    r = client.post(f"/api/plm/revisions/{rid}/documents", json={
        "doc_type": "sds", "properties": {"flash_point_c": 30, "initial_boiling_point_c": 138}})
    assert r.status_code == 201, r.text
    first = r.json()
    assert first["version_no"] == 1 and first["status"] == "draft" and first["has_pdf"] and first["has_docx"]
    oel = first["content"]["sections"][7]["tables"][0]["rows"][0]
    assert oel[:2] == ["二甲苯", "1330-20-7"] and oel[3] == 50
    listing = first["content"]["sections"][14]["fields"][1]["value"]
    assert "序号 2289" in listing[0]
    assert first["content"]["meta"]["classification_id"]
    pdf = client.get(f"/api/plm/documents/{first['id']}/file/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert client.get(f"/api/plm/documents/{first['id']}/file/docx").content.startswith(b"PK")

    second = client.post(f"/api/plm/revisions/{rid}/documents", json={
        "doc_type": "sds", "properties": {"flash_point_c": 20, "initial_boiling_point_c": 138}}).json()
    assert second["version_no"] == 2 and second["supersedes_id"] == first["id"]
    diff = client.get(f"/api/plm/documents/{second['id']}/diff").json()
    assert diff["major"] and diff["from"] == first["id"]
    assert [d["version_no"] for d in client.get(f"/api/plm/revisions/{rid}/documents").json()] == [1, 2]


def test_label_tds_and_coa(api):
    client, pid, rid = api
    lab = client.post(f"/api/plm/revisions/{rid}/documents", json={
        "doc_type": "label", "properties": {"flash_point_c": 30, "initial_boiling_point_c": 138},
        "label": {"container_volume_l": 18, "size_mm": [100, 150]}}).json()
    assert lab["has_pdf"] and not lab["has_docx"]
    tds = client.post(f"/api/plm/revisions/{rid}/documents", json={
        "doc_type": "tds", "lang": "en-US", "product": {"name_en": "FX-100 Primer"},
        "tds": {"product": {"description": "Two-pack epoxy primer"}, "application": {"method": "spray"},
                "properties": [{"key": "density_g_cm3", "label": "Density", "value": 1.35,
                                "unit": "g/cm³", "method": "GB/T 6750"}]}}).json()
    assert tds["content"]["title"] == "Technical Data Sheet" and tds["missing"] == []

    bid = client.post(f"/api/plm/products/{pid}/batches", json={"batch_no": "B2610-01", "produced_at": "2026-10-15"}).json()["id"]
    client.post(f"/api/plm/batches/{bid}/qc-results", json={"results": [{"metric": "solids", "value": 57.1},
                                                                       {"metric": "visc", "value": 95}]})
    coa = client.post(f"/api/plm/revisions/{rid}/documents", json={"doc_type": "coa", "coa": {"batch_id": bid}}).json()
    assert coa["content"]["meta"]["overall"] == "FAIL" and coa["content"]["meta"]["verdicts"] == {"solids": "PASS", "visc": "FAIL"}
    assert coa["batch_id"] == bid
    assert client.post(f"/api/plm/revisions/{rid}/documents", json={"doc_type": "coa", "coa": {}}).status_code == 422


def test_tds_features_and_applications_split_string_and_list():
    """P1-22: a UI text-area string must split, not become one-character items."""

    def _features(doc):
        fields = doc["sections"][0]["fields"]
        return {f["key"]: f["value"] for f in fields}

    doc = build_tds(TdsContext(product={**PRODUCT, "description": "底漆",
                                       "features": "耐腐蚀、耐磨，耐高温；耐酸",
                                       "applications": "钢结构、桥梁"},
                               company=COMPANY, properties=[],
                               application={"method": "喷涂"}))
    got = _features(doc)
    assert got["features"] == ["耐腐蚀", "耐磨", "耐高温", "耐酸"]
    assert got["applications"] == ["钢结构", "桥梁"]
    doc = build_tds(TdsContext(product={**PRODUCT, "description": "底漆",
                                       "features": ["耐腐蚀", "耐磨"],
                                       "applications": ["钢结构"]},
                               company=COMPANY, properties=[],
                               application={"method": "喷涂"}))
    got = _features(doc)
    assert got["features"] == ["耐腐蚀", "耐磨"] and got["applications"] == ["钢结构"]
    assert split_list_field("", "features") == [] and split_list_field(None, "features") == []
    with pytest.raises(ValueError, match="字符串列表"):
        split_list_field(42, "features")
    with pytest.raises(ValueError, match="非字符串项"):
        split_list_field(["耐腐蚀", 7], "features")


def test_tds_string_features_via_api_are_split_not_characters(api):
    """P1-22 API path: string features from a caller must not reach build_tds raw."""
    client, pid, rid = api
    r = client.post(f"/api/plm/revisions/{rid}/documents", json={
        "doc_type": "tds", "tds": {"product": {"description": "双组分环氧底漆", "features": "耐腐蚀、耐磨",
                                              "applications": "钢结构，桥梁"},
                                   "properties": [{"key": "density_g_cm3", "label": "密度", "value": 1.35, "unit": "g/cm³"}],
                                   "application": {"method": "喷涂"}}})
    assert r.status_code == 201, r.text
    fields = {f["key"]: f["value"] for f in r.json()["content"]["sections"][0]["fields"]}
    assert fields["features"] == ["耐腐蚀", "耐磨"]
    assert fields["applications"] == ["钢结构", "桥梁"]
    bad = client.post(f"/api/plm/revisions/{rid}/documents", json={
        "doc_type": "tds", "tds": {"product": {"description": "双组分环氧底漆", "features": 42},
                                   "properties": [{"key": "density_g_cm3", "label": "密度", "value": 1.35, "unit": "g/cm³"}],
                                   "application": {"method": "喷涂"}}})
    assert bad.status_code == 422, bad.text


def test_coa_two_batches_each_get_version_one(api):
    """P1-20 regression: the second batch's first COA used to 500 on the
    (revision, doc_type, lang, version_no) unique constraint."""
    client, pid, rid = api
    ids = []
    for batch_no in ("B2610-01", "B2610-02"):
        bid = client.post(f"/api/plm/products/{pid}/batches", json={"batch_no": batch_no}).json()["id"]
        r = client.post(f"/api/plm/revisions/{rid}/documents", json={"doc_type": "coa", "coa": {"batch_id": bid}})
        assert r.status_code == 201, r.text
        body = r.json()
        assert body["version_no"] == 1 and body["batch_id"] == bid
        ids.append(body["id"])
    assert ids[0] != ids[1]
    # A second COA for the first batch still increments within that batch.
    docs = client.get(f"/api/plm/revisions/{rid}/documents").json()
    first_batch = next(d for d in docs if d["doc_type"] == "coa" and d["batch_id"] and d["version_no"] == 1)
    again = client.post(f"/api/plm/revisions/{rid}/documents",
                        json={"doc_type": "coa", "coa": {"batch_id": first_batch["batch_id"]}})
    assert again.status_code == 201, again.text
    assert again.json()["version_no"] == 2 and again.json()["batch_id"] == first_batch["batch_id"]


def test_frozen_revision_and_errors(api):
    client, pid, rid = api
    from app.db.database import default_session_factory
    from app.db.models_plm import ProductRevisionRow

    assert client.post("/api/plm/products", json={"code": "FX-100", "name_zh": "dup"}).status_code == 422
    with default_session_factory()() as s:
        s.get(ProductRevisionRow, rid).frozen_at = datetime(2026, 10, 16)
        s.commit()
    r = client.post(f"/api/plm/revisions/{rid}/documents", json={"doc_type": "sds"})
    assert r.status_code == 422 and "冻结" in r.json()["detail"]
    nofv = client.post(f"/api/plm/products/{pid}/revisions", json={}).json()["id"]
    assert client.post(f"/api/plm/revisions/{nofv}/documents", json={"doc_type": "sds"}).status_code == 422
    assert client.get("/api/plm/documents/nope").status_code == 404


def test_template_upload_is_used_for_word(api):
    from docx import Document

    client, pid, rid = api
    tpl = Document()
    tpl.add_paragraph("企业模板：{{ header.product_name }}")
    buf = io.BytesIO()
    tpl.save(buf)
    r = client.post("/api/plm/templates", data={"doc_type": "sds", "lang": "zh-CN", "name": "公司版式"},
                    files={"file": ("sds.docx", buf.getvalue())})
    assert r.status_code == 201, r.text
    assert client.post("/api/plm/templates", data={"doc_type": "sds", "name": "bad"},
                       files={"file": ("x.docx", b"not a docx")}).status_code == 422
    doc = client.post(f"/api/plm/revisions/{rid}/documents", json={
        "doc_type": "sds", "properties": {"flash_point_c": 30, "initial_boiling_point_c": 138}}).json()
    import docx

    word = docx.Document(io.BytesIO(client.get(f"/api/plm/documents/{doc['id']}/file/docx").content))
    assert word.paragraphs[0].text == "企业模板：FX-100 环氧底漆"


def test_2_8_coa_prefers_test_method_standard():
    """补齐计划 2-8：COA 检验方法列优先显示 test_method_standard，空则回退 method（老数据）。"""
    specs = [SpecItem("solids", "固含量", "%", "老方法", 55, 60, test_method_standard="GB/T 1725"),
             SpecItem("visc", "黏度", "s", "GB/T 1723", max=90)]
    doc = build_coa(CoaContext(PRODUCT, COMPANY, {"batch_no": "B1"}, specs,
                               [ResultItem("solids", 57.2), ResultItem("visc", 80)]))
    tables = doc["sections"][1]["tables"]
    results = next(t for t in tables if t["key"] == "results")["rows"]
    by_label = {r[0]: r for r in results}
    assert by_label["固含量"][5] == "GB/T 1725"
    assert by_label["黏度"][5] == "GB/T 1723"


def test_doc_payload_debug_log_flag_gated(api, monkeypatch, caplog):
    """FORMUNEX_PLM_DEBUG_DOC_PAYLOAD: on → payload fields logged, off → silent.
    Trade secrets must never appear in the log line."""
    import logging

    client, pid, rid = api
    body = {"doc_type": "tds",
            "tds": {"product": {"description": "双组分环氧底漆"},
                    "properties": [{"key": "density_g_cm3", "label": "密度", "value": 1.35, "unit": "g/cm³"}]},
            "trade_secrets": {"1330-20-7": "芳烃溶剂"}}
    with caplog.at_level(logging.INFO, logger="app.api.plm"):
        monkeypatch.setenv("FORMUNEX_PLM_DEBUG_DOC_PAYLOAD", "true")
        r = client.post(f"/api/plm/revisions/{rid}/documents", json=body)
        assert r.status_code == 201, r.text
    line = next((rec.message for rec in caplog.records if "plm doc payload" in rec.message), None)
    assert line is not None
    assert "density_g_cm3" in line and "tds.properties=" in line
    assert "芳烃溶剂" not in line and "trade_secrets" not in line
    caplog.clear()
    with caplog.at_level(logging.INFO, logger="app.api.plm"):
        monkeypatch.delenv("FORMUNEX_PLM_DEBUG_DOC_PAYLOAD", raising=False)
        r = client.post(f"/api/plm/revisions/{rid}/documents", json=body)
        assert r.status_code == 201, r.text
    assert not any("plm doc payload" in rec.message for rec in caplog.records)
