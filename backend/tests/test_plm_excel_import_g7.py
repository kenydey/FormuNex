"""G7 Excel import: formulations + companies (materials already covered)."""
from __future__ import annotations

import io
import json

import pytest

from app.config import get_settings
from app.db.database import Base, make_engine, make_session_factory
from app.db.models_plm import CompanyRow, FormulaMappingRow, ProductRow
from app.services import material_io
from app.services.plm import excel_import as xi


@pytest.fixture()
def sf(tmp_path, monkeypatch):
    monkeypatch.setenv("FORMUNEX_PLM_STORAGE_DIR", str(tmp_path / "plm"))
    get_settings.cache_clear()
    engine = make_engine(f"sqlite:///{tmp_path}/g7.db")
    Base.metadata.create_all(engine)
    yield make_session_factory(engine)
    engine.dispose()
    get_settings.cache_clear()


def _xlsx(headers, rows) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parsed(headers, rows):
    return material_io.parse_xlsx_bytes(_xlsx(headers, rows))


FORMULA_HEADERS = [h for h, _ in xi.FORMULA_HEADERS]
COMPANY_HEADERS = [h for h, _ in xi.COMPANY_HEADERS]
COMPS = json.dumps([{"name": "水", "pct": 60, "cas": "7732-18-5"},
                    {"name": "测试组分", "pct": 40}], ensure_ascii=False)


# --------------------------------------------------------------------------- mapping


def test_map_rows_headers_aliases_and_row_numbers():
    rows = _parsed(["产品代码", "产品 名称", "code", "未知列", "成分列表"],
                   [["P-1", "磷化剂", "IGNORED", "x", COMPS],
                    [None, None, None, None, None]])
    mapped = xi.map_rows(rows, xi.FORMULA_HEADERS, xi.FORMULA_ALIASES)
    assert len(mapped) == 1
    row_no, rec = mapped[0]
    assert row_no == 2  # header = row 1
    # The first matching header wins: 产品代码 beats the later "code" alias.
    assert rec["product_code"] == "P-1" and rec["product_name_zh"] == "磷化剂"
    assert rec["components_json"] == COMPS and "未知列" not in str(rec)


# --------------------------------------------------------------------------- formulations


def test_formula_import_created_skipped_errors(sf):
    rows = _parsed(FORMULA_HEADERS, [
        ["P-1", "磷化剂", "", "测试公司", "010-12345678", 1, "SRC-1", COMPS, ""],
        ["P-1", "磷化剂", "", "测试公司", "010-12345678", 1, "SRC-1", COMPS, ""],
        ["P-2", "坏配方", "", "", "", 1, "SRC-2", "not json", ""],
        ["", "无代码", "", "", "", 1, "SRC-3", COMPS, ""],
    ])
    mapped = xi.map_rows(rows, xi.FORMULA_HEADERS, xi.FORMULA_ALIASES)
    with sf() as s:
        report = xi.import_formula_rows(s, mapped)
        assert len(report["created"]) == 1 and len(report["skipped"]) == 1
        assert len(report["errors"]) == 2
        assert report["errors"][0]["row"] == 4  # bad JSON, spreadsheet row
        assert report["errors"][1]["error"].startswith("配方记录缺少产品代码")
        assert s.query(ProductRow).count() == 1
        assert s.query(FormulaMappingRow).count() == 1
        assert s.query(CompanyRow).count() == 1


def test_formula_import_dry_run_persists_nothing(sf):
    rows = _parsed(FORMULA_HEADERS, [["P-9", "干跑", "", "", "", 1, "SRC-9", COMPS, ""]])
    mapped = xi.map_rows(rows, xi.FORMULA_HEADERS, xi.FORMULA_ALIASES)
    with sf() as s:
        report = xi.import_formula_rows(s, mapped, dry_run=True)
        assert report["dry_run"] is True and len(report["created"]) == 1
        assert s.query(ProductRow).count() == 0


def test_formula_import_company_without_emergency_warns(sf):
    rows = _parsed(FORMULA_HEADERS, [["P-3", "无应急", "", "缺电话公司", "", 1, "SRC-3", COMPS, ""]])
    mapped = xi.map_rows(rows, xi.FORMULA_HEADERS, xi.FORMULA_ALIASES)
    with sf() as s:
        report = xi.import_formula_rows(s, mapped)
        assert len(report["created"]) == 1 and report["warnings"]
        assert s.query(CompanyRow).count() == 0  # no placeholder company


# --------------------------------------------------------------------------- companies


def test_company_import_created_skipped_errors(sf):
    rows = _parsed(COMPANY_HEADERS, [
        ["甲公司", "A Co", "北京", "010-1", "a@x.com", "010-999", "91110000"],
        ["甲公司", "A Co", "", "", "", "010-999", ""],
        ["乙公司", "", "", "", "", "", ""],
    ])
    mapped = xi.map_rows(rows, xi.COMPANY_HEADERS, xi.COMPANY_ALIASES)
    with sf() as s:
        report = xi.import_company_rows(s, mapped)
        assert len(report["created"]) == 1 and len(report["skipped"]) == 1
        assert len(report["errors"]) == 1 and "应急电话" in report["errors"][0]["error"]
        assert s.query(CompanyRow).count() == 1


def test_company_import_dry_run_and_template(sf):
    rows = _parsed(COMPANY_HEADERS, [["丙公司", "", "", "", "", "010-999", ""]])
    mapped = xi.map_rows(rows, xi.COMPANY_HEADERS, xi.COMPANY_ALIASES)
    with sf() as s:
        report = xi.import_company_rows(s, mapped, dry_run=True)
        assert len(report["created"]) == 1 and s.query(CompanyRow).count() == 0
    payload = xi.build_template(xi.COMPANY_HEADERS)
    parsed = material_io.parse_xlsx_bytes(payload)
    assert parsed == []  # header-only: no data rows
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True)
    assert [c.value for c in next(wb.active.iter_rows())] == COMPANY_HEADERS
